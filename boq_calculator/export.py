"""Export the current calculator state to an Excel BOQ workbook."""

from __future__ import annotations

import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font

from .models import OnlineEvent, OfflineEvent, StreamingSetup
from .calculations import Totals

BOLD = Font(bold=True)


def _write_rows(ws, rows):
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = BOLD


def build_workbook(
    setup: StreamingSetup,
    online_events: List[OnlineEvent],
    offline_events: List[OfflineEvent],
    totals: Totals,
) -> Workbook:
    wb = Workbook()

    ws_req = wb.active
    ws_req.title = "Client Requirements"
    _write_rows(
        ws_req,
        [
            ["Streaming Requirements", ""],
            ["Total Camera", setup.total_cameras],
            ["Camera Operational Hours", setup.operational_hours],
            ["DVR Storage Days", setup.dvr_days],
            ["NDVR Storage Hours", setup.ndvr_hours],
            ["Incoming Bandwidth mbps / Camera", setup.bandwidth_mbps_per_camera],
        ],
    )

    ws_out = wb.create_sheet("Outcomes")
    _write_rows(
        ws_out,
        [
            ["Metric", "Value"],
            ["Total GPU Requirement (GB)", totals.gpu_gb],
            ["GPU Cards (16GB equivalent)", totals.gpu_cards],
            ["Total RAM Requirement (GB)", totals.ram_gb],
            ["Total vCPU Requirement", totals.vcpu],
            ["Total Physical Cores Required", totals.physical_cores],
            ["Server Quantity", totals.servers],
            ["DVR Storage (TB)", totals.dvr_tb],
            ["NDVR Storage (TB)", totals.ndvr_tb],
            ["Incoming Bandwidth (Mbps)", totals.bandwidth_mbps],
        ],
    )

    ws_online = wb.create_sheet("Online-Event")
    rows = [[
        "Event", "Cameras", "Model / Stage", "FPS Load / Camera",
        "Capacity (FPS/Instance)", "Instances", "GPU MB/Instance",
        "GPU Required MB", "RAM MB/Instance", "RAM Required MB",
        "vCPU/Instance", "Total vCPU",
    ]]
    for ev in online_events:
        if not ev.enabled:
            continue
        for s in ev.stages:
            r = s.resource_usage(ev.cameras)
            rows.append([
                ev.name, ev.cameras, s.name, s.mult, s.cap,
                r["instances"], s.gpu_mb, r["gpu_mb"], s.ram_mb, r["ram_mb"],
                s.vcpu, r["vcpu"],
            ])
    _write_rows(ws_online, rows)

    ws_offline = wb.create_sheet("Offline-Event")
    rows = [[
        "Event", "Recordings/Day", "Min/Recording", "Capacity/Day/Instance",
        "Instances", "GPU MB/Instance", "GPU Required MB", "RAM MB/Instance",
        "RAM Required MB", "vCPU/Instance", "Total vCPU",
    ]]
    for ev in offline_events:
        if not ev.enabled:
            continue
        r = ev.resource_usage()
        rows.append([
            ev.name, ev.recordings_per_day, ev.minutes_per_recording,
            round(ev.capacity_per_day(), 2), r["instances"], ev.gpu_mb,
            r["gpu_mb"], ev.ram_mb, r["ram_mb"], ev.vcpu, r["vcpu"],
        ])
    _write_rows(ws_offline, rows)

    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
